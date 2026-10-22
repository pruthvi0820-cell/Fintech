"""Bytes → plain text, the way a basic ATS parser sees a resume.

The format comes from the file's first bytes, never its extension: a renamed file still parses
and a mislabelled one is refused. Size, page and zip limits stop oversized or hostile uploads.
"""

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from .findings import HIGH, MEDIUM, Finding

MAX_BYTES = 5 * 1024 * 1024            # real resumes are well under 1 MB
MAX_PAGES = 10
MAX_UNZIPPED_BYTES = 50 * 1024 * 1024  # DOCX is a zip: refuse zip bombs before python-docx opens it
MIN_CHARS_PER_PAGE = 100               # below this a PDF page is almost certainly an image (scan/export)


_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x80-\x9f]")


class ExtractionError(ValueError):
    """The file cannot be read as a resume at all (wrong type, too big, encrypted, corrupt)."""


@dataclass(frozen=True)
class Extracted:
    text: str                       # body text in reading order
    source_format: str              # "pdf" | "docx" | "txt"
    pages: int | None = None        # PDF only
    header_footer_text: str = ""    # DOCX only: many ATS parsers skip these areas
    findings: tuple[Finding, ...] = field(default_factory=tuple)


def extract_file(path: str | Path) -> Extracted:
    p = Path(path)
    try:
        size = p.stat().st_size
    except OSError as e:
        raise ExtractionError(f"cannot read {p.name}: {e.strerror or e}") from e
    if size > MAX_BYTES:
        raise ExtractionError(f"{p.name} is {size / 1e6:.1f} MB; the limit is {MAX_BYTES / 1e6:.0f} MB")
    return extract_bytes(p.read_bytes(), p.name)


def extract_bytes(data: bytes, filename: str = "resume") -> Extracted:
    if not data:
        raise ExtractionError(f"{filename} is empty")
    if len(data) > MAX_BYTES:
        raise ExtractionError(f"{filename} is over the {MAX_BYTES / 1e6:.0f} MB limit")
    kind = sniff(data)
    if kind == "pdf":
        return _pdf(data)
    if kind == "docx":
        return _docx(data)
    if kind == "txt":
        return Extracted(text=normalize(data.decode("utf-8")), source_format="txt")
    raise ExtractionError(f"{filename} is not a PDF, DOCX or UTF-8 text file")


def sniff(data: bytes) -> str | None:
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                return "docx" if "word/document.xml" in z.namelist() else None
        except zipfile.BadZipFile:
            return None
    if data.startswith(b"\xd0\xcf\x11\xe0"):
        return None                     # legacy .doc (OLE): not supported, caller gets a clear error
    try:
        data.decode("utf-8")
    except UnicodeDecodeError:
        return None
    return "txt" if b"\x00" not in data else None


def normalize(text: str) -> str:
    """NFKC folds ligatures (ﬁ → fi) and odd spaces; every bullet glyph becomes '•'."""
    text = unicodedata.normalize("NFKC", text).replace("\r\n", "\n").replace("\r", "\n")
    # "\x7f" is what some resume builders' bullet glyph decodes to; other control characters are noise
    for glyph in ("●", "▪", "■", "•", "‣", "⁃", "", "", "➢", "\x7f"):
        text = text.replace(glyph, "•")
    text = _CONTROL.sub("", text)
    lines = (" ".join(line.split()) for line in text.split("\n"))
    out, blank = [], False
    for line in lines:                  # keep single blank lines (they separate roles), drop runs of them
        if line or not blank:
            if out and out[-1] == "•" and line:          # a bullet alone on its line belongs to the next one
                out[-1] = f"• {line}"
            else:
                out.append(line)
        blank = not line
    return "\n".join(out).strip()


def _pdf(data: bytes) -> Extracted:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    findings: list[Finding] = []
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractionError("the PDF is password-protected; ATS parsers cannot open it either")
        page_count = len(reader.pages)
    except ExtractionError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as e:
        raise ExtractionError(f"the PDF is damaged or not a real PDF ({type(e).__name__})") from e
    except Exception as e:  # pypdf raises assorted errors on broken files; never let one escape as a crash
        raise ExtractionError(f"the PDF could not be opened ({type(e).__name__})") from e

    if page_count == 0:
        raise ExtractionError("the PDF has no pages")
    if page_count > MAX_PAGES:
        findings.append(Finding("too_many_pages", MEDIUM,
                                f"{page_count} pages; only the first {MAX_PAGES} were read"))

    texts, image_pages = [], 0
    for number, page in enumerate(reader.pages[:MAX_PAGES], start=1):
        try:
            t = page.extract_text() or ""
        except Exception as e:  # one bad page should not lose the rest of the resume
            findings.append(Finding("pdf_page_unreadable", HIGH,
                                    f"page {number} could not be read ({type(e).__name__})"))
            t = ""
        if len(t.strip()) < MIN_CHARS_PER_PAGE:
            image_pages += 1
        texts.append(t)

    text = normalize("\n".join(texts))
    read = min(page_count, MAX_PAGES)
    if not text:
        findings.append(Finding("pdf_no_text", HIGH,
                                "no selectable text: the PDF is an image (scan or design export). "
                                "An ATS sees a blank resume. Export it again as a text PDF."))
    elif image_pages:
        findings.append(Finding("pdf_low_text", MEDIUM,
                                f"{image_pages} of {read} pages have almost no text; parts may be images"))
    return Extracted(text=text, source_format="pdf", pages=page_count, findings=tuple(findings))


def _docx(data: bytes) -> Extracted:
    from docx import Document

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            unzipped = sum(i.file_size for i in z.infolist())
    except zipfile.BadZipFile as e:
        raise ExtractionError("the DOCX is damaged") from e
    if unzipped > MAX_UNZIPPED_BYTES:
        raise ExtractionError("the DOCX expands to an unreasonable size; refusing to open it")

    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:  # python-docx surfaces lxml/zip/KeyError variants for broken files
        raise ExtractionError(f"the DOCX could not be opened ({type(e).__name__})") from e

    findings: list[Finding] = []
    body = [p.text for p in doc.paragraphs]
    if doc.tables:
        findings.append(Finding("docx_tables", MEDIUM,
                                f"{len(doc.tables)} table(s): many ATS parsers read table cells out of order"))
        for table in doc.tables:
            for row in table.rows:
                cells = []
                for cell in row.cells:   # merged cells repeat the same object; keep each once
                    if not cells or cell.text != cells[-1]:
                        cells.append(cell.text)
                body.append(" | ".join(c for c in cells if c.strip()))

    edges = []
    for section in doc.sections:
        for part in (section.header, section.footer):
            edges.extend(p.text for p in part.paragraphs)
    edge_text = normalize("\n".join(edges))

    text = normalize("\n".join(body))
    if not text:
        findings.append(Finding("docx_no_text", HIGH, "no body text found; content may be in text boxes or images"))
    return Extracted(text=text, source_format="docx", header_footer_text=edge_text, findings=tuple(findings))
