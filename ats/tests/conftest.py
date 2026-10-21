"""Builders for real PDF/DOCX bytes, so tests exercise the same code path as uploads."""

import io
from datetime import date

import pytest

TODAY = date(2026, 10, 21)

SAMPLE = """\
PRIYA SHARMA
Bengaluru, India | priya.sharma@example.com | +91 98765 43210
linkedin.com/in/priya-sharma | github.com/priyasharma

PROFESSIONAL SUMMARY
Backend engineer with 3 years building payment APIs in Python.

WORK EXPERIENCE
Software Engineer | Razorpay
Jan 2024 – Present
• Built a reconciliation service in FastAPI handling 2M transactions a day.
• Cut settlement latency by 35% by moving batch jobs to Kafka consumers.
Associate Software Engineer, Infosys          Jul 2022 - Dec 2023
• Maintained Java Spring microservices for a banking client.

EDUCATION
B.Tech, Computer Science, VTU    2018 – 2022

Skills: Python, FastAPI, PostgreSQL, Kafka, Docker, AWS
"""


def _pdf_escape(s: str) -> str:
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(pages: list[list[str]]) -> bytes:
    """Minimal valid PDF: one Helvetica text block per page, correct xref offsets."""
    n = len(pages)
    objs: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [" + " ".join(f"{4 + 2 * i} 0 R" for i in range(n)).encode() + b"] /Count "
        + str(n).encode() + b" >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    for i, lines in enumerate(pages):
        ops = "BT /F1 10 Tf 14 TL 50 800 Td " + " ".join(f"({_pdf_escape(ln)}) Tj T*" for ln in lines) + " ET"
        stream = ops.encode("cp1252")
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 3 0 R >> >> "
                    f"/Contents {5 + 2 * i} 0 R >>".encode())
        objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for num, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{num} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None, header: str | None = None) -> bytes:
    from docx import Document

    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    if header is not None:
        doc.sections[0].header.paragraphs[0].text = header
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def sample_text() -> str:
    return SAMPLE
